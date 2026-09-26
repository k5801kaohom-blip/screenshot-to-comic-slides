#!/usr/bin/env python3
"""Fallback: package ordered full-slide images from a completed Slides image deck into PPTX."""

import argparse
import io
import re
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def natural_key(path: Path):
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", path.name)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input_dir", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--pattern", default="slide_*")
    ap.add_argument("--expected-count", type=int)
    args = ap.parse_args()

    images = sorted(
        [p for p in args.input_dir.glob(args.pattern) if p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}],
        key=natural_key
    )
    if not images:
        raise SystemExit("No slide images found")
    if args.expected_count is not None and len(images) != args.expected_count:
        raise SystemExit(f"Expected {args.expected_count} images, found {len(images)}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for path in images:
        im = Image.open(path).convert("RGB")
        if im.width * 9 != im.height * 16:
            canvas = Image.new("RGB", (1280, 720), "white")
            im.thumbnail((1280, 720), Image.Resampling.LANCZOS)
            canvas.paste(im, ((1280 - im.width) // 2, (720 - im.height) // 2))
            im = canvas
        buffer = io.BytesIO()
        im.save(buffer, format="PNG")
        buffer.seek(0)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(buffer, 0, 0, width=prs.slide_width, height=prs.slide_height)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
